import csv
import hashlib
import io
import json
from datetime import datetime

from openpyxl import load_workbook

from collector.base import RawItem

_SENTIMENTS = {"positive": "positive", "positivo": "positive", "negative": "negative", "negativo": "negative", "neutral": "neutral", "neutro": "neutral"}
_STANCES = {
    "in_favor": "in_favor",
    "a_favor": "in_favor",
    "against": "against",
    "en_contra": "against",
    "mixed": "mixed",
    "mixto": "mixed",
    "not_applicable": "not_applicable",
    "no_aplica": "not_applicable",
}


def _parse_dt(value: str | None) -> datetime | None:
    if not value:
        return None
    text = value.strip().replace("Z", "+00:00")
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def _int(value: str | int | None) -> int:
    if value in (None, ""):
        return 0
    return int(value)


def stable_external_id(row: dict, text: str) -> str:
    raw = row.get("id")
    if raw in (None, ""):
        raw = row.get("external_id")
    explicit = "" if raw in (None, "") else str(raw).strip()
    if explicit:
        return explicit[:120]
    digest = hashlib.sha256(text.casefold().encode("utf-8")).hexdigest()
    return f"manual-{digest[:32]}"


def _target_hint(row: dict) -> str | None:
    raw = row.get("objetivo") or row.get("target") or row.get("objetivo_nombre")
    if raw in (None, ""):
        return None
    text = str(raw).strip()
    return text or None


def _row_to_item(row: dict) -> RawItem | None:
    text = (row.get("texto") or row.get("text") or "").strip()
    if not text:
        return None
    sentiment_raw = (row.get("sentimiento") or row.get("sentiment") or "").strip().casefold()
    stance_raw = (row.get("postura") or row.get("stance") or "").strip().casefold()
    confidence = row.get("confianza") or row.get("confidence")
    relevance = row.get("relevancia") or row.get("relevance")
    return RawItem(
        source=(row.get("fuente") or row.get("source") or "manual_upload").strip() or "manual_upload",
        external_id=stable_external_id(row, text),
        text=text,
        url=(row.get("url") or None),
        author_handle=(row.get("autor") or row.get("author") or None),
        published_at=_parse_dt(row.get("fecha") or row.get("published_at")),
        likes=_int(row.get("me_gusta") or row.get("likes")),
        replies=_int(row.get("respuestas") or row.get("replies")),
        shares=_int(row.get("compartidos") or row.get("shares")),
        views=_int(row.get("vistas") or row.get("views")),
        author_followers=_int(row.get("seguidores") or row.get("followers")),
        geo_state=row.get("estado") or row.get("geo_state"),
        geo_municipality=row.get("municipio") or row.get("geo_municipality"),
        sentiment=_SENTIMENTS.get(sentiment_raw),
        stance=_STANCES.get(stance_raw),
        theme=row.get("tema") or row.get("theme"),
        confidence=float(confidence) if confidence not in (None, "") else None,
        relevance=float(relevance) if relevance not in (None, "") else None,
        target_hint=_target_hint(row),
        license_note="Carga del analista. El estudio declara derecho de uso.",
    )


def parse_csv(content: bytes) -> list[RawItem]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    items = []
    for row in reader:
        item = _row_to_item(row)
        if item:
            items.append(item)
    return items


def parse_json(content: bytes) -> list[RawItem]:
    payload = json.loads(content.decode("utf-8"))
    rows = payload if isinstance(payload, list) else payload.get("mentions", [])
    items = []
    for row in rows:
        item = _row_to_item(row)
        if item:
            items.append(item)
    return items


def parse_xlsx(content: bytes) -> list[RawItem]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    items = []
    for values in rows:
        row = {headers[pos]: "" if value is None else str(value) for pos, value in enumerate(values)}
        item = _row_to_item(row)
        if item:
            items.append(item)
    return items


def parse_upload(content: bytes, filename: str) -> list[RawItem]:
    name = filename.casefold()
    if name.endswith(".json"):
        return parse_json(content)
    if name.endswith(".xlsx"):
        return parse_xlsx(content)
    return parse_csv(content)
